"""運行環境模式的 Python 節點（ADR 0020）：runtime 階段、rpa_script、共用工作資料夾、斷言與失敗。

運行環境平常用一個假的：標準函式庫的 venv（不裝 pip）加上 rpa-runtime.json，以 RPA_PYTHON_RUNTIME
指給 Runner；rpa_script 和真的運行環境一樣由 Runner 提供（pyruntime.script_library）。
CI 建好真的運行環境後（python-runtime.yml）以 ``RPA_TEST_PYTHON_RUNTIME`` 指給這些測試，每個系統
（含 Windows）用真的再跑一次，另外加上用到運行環境套件的測試。不需要瀏覽器。
"""

import json
import os
import subprocess
import sys
import time
from collections.abc import Iterator
from pathlib import Path
from typing import cast

import pytest

from rpa_core.i18n import use_locale
from rpa_runner.files import FACTS, OUTPUT, RESULT, STEPS
from rpa_runner.pyruntime import ENV_RUNTIME, MANIFEST, platform_tag, read_manifest
from rpa_runner.run import RunResult, execute

REAL_RUNTIME: str = os.environ.get("RPA_TEST_PYTHON_RUNTIME", "")
"""已經解壓的真的運行環境（CI 的 python-runtime.yml 設定）；沒設定時用假的。"""

pytestmark = pytest.mark.skipif(
    os.name == "nt" and not REAL_RUNTIME, reason="假的運行環境用 venv 的 bin/python3"
)
real_only = pytest.mark.skipif(not REAL_RUNTIME, reason="要用真的運行環境")


@pytest.fixture(scope="module")
def runtime_root(tmp_path_factory: pytest.TempPathFactory) -> Path:
    if REAL_RUNTIME:
        return Path(REAL_RUNTIME)
    root = tmp_path_factory.mktemp("runtime")
    subprocess.run([sys.executable, "-m", "venv", "--without-pip", str(root)], check=True)
    manifest = {"version": "1.0.0", "platform": platform_tag(), "python": sys.version.split()[0]}
    (root / MANIFEST).write_text(json.dumps(manifest), encoding="utf-8")
    return root


@pytest.fixture(scope="module")
def cache(tmp_path_factory: pytest.TempPathFactory) -> Path:
    return tmp_path_factory.mktemp("cache")


@pytest.fixture(autouse=True)
def _runtime(runtime_root: Path, cache: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    monkeypatch.setenv(ENV_RUNTIME, str(runtime_root))
    monkeypatch.setenv("RPA_CACHE_DIR", str(cache))  # Runner 的 rpa_script 放在這裡
    with use_locale("zh-Hant"):
        yield


def node(
    folder: Path,
    automation: str,
    *,
    extra: str = "",
    files: dict[str, str] | None = None,
    name: str = "node",
    assertion: str | None = None,
) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "automation.py").write_text(automation, encoding="utf-8")
    if assertion is not None:
        (folder / "assertion.py").write_text(assertion, encoding="utf-8")
    for relative, content in (files or {}).items():
        (folder / relative).parent.mkdir(parents=True, exist_ok=True)
        (folder / relative).write_text(content, encoding="utf-8")
    path = folder / "scenario.yaml"
    path.write_text(
        f"schemaVersion: 1\nid: {name}\nname: 節點\n"
        + extra
        + "script:\n  runtime: '1'\n  automation: automation.py\n"
        + ("  assertion: assertion.py\n" if assertion is not None else "")
        + ("  files: [" + ", ".join(files) + "]\n" if files else ""),
        encoding="utf-8",
    )
    return path


def read(result: RunResult, name: str) -> object:
    return json.loads((result.directory / name).read_text(encoding="utf-8"))


def steps_of(result: RunResult) -> list[dict[str, object]]:
    return cast("list[dict[str, object]]", read(result, STEPS))


MERGE = """\
import os
from rpa_script import Run, StepError, automation
from helpers import total


@automation
def main(run: Run) -> None:
    run.log("token " + os.environ["API_TOKEN"])
    with run.step("write", "寫出"):
        (run.output_dir / "summary.csv").write_text("a,b\\n", encoding="utf-8")
        run.preview("資料", [{"name": n, "value": v} for n, v in run.inputs["rows"]])
    for name, value in run.inputs["rows"]:
        with run.step("check", f"檢查 {name}"):
            if value < 0:
                raise StepError(f"{name} 是負數")
    run.facts.update(
        total=total(run.inputs["rows"]),
        byName={n: v for n, v in run.inputs["rows"]},
        cwd=os.getcwd(),
        path=os.environ.get("PYTHONPATH", ""),
    )
"""
MERGE_EXTRA = (
    "inputs:\n  rows: { type: array, default: [[A, 3], [B, 4]] }\n"
    "outputs:\n"
    "  total: { type: number, from: '{{ facts.total }}' }\n"
    "  byName: { type: object, from: '{{ facts.byName }}' }\n"
    "secrets: [API_TOKEN]\n"
    "verify:\n  - { value: '{{ facts.total }}', matches: '^[1-9]' }\n"
)


def test_runtime_node_end_to_end(
    tmp_path: Path, runtime_root: Path, cache: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("API_TOKEN", "secret-value-123")
    monkeypatch.setenv("PYTHONPATH", "/should/not/leak")
    path = node(
        tmp_path / "scene",
        MERGE,
        extra=MERGE_EXTRA,
        files={"helpers.py": "def total(rows):\n    return sum(v for _, v in rows)\n"},
    )
    result = execute(path, {}, out_root=tmp_path / "runs")

    assert result.status == "passed", (result.error, result.problems)
    assert result.output == {"total": 7, "byName": {"A": 3, "B": 4}}
    assert [stage.name for stage in result.stages] == [
        "input",
        "runtime",
        "automation",
        "verify",
        "output",
    ]
    assert result.python is not None
    assert result.python["mode"] == "runtime"
    assert result.python["runtime"] == read_manifest(runtime_root)["version"]
    # rpa_script 是 Runner 提供的
    assert list(cache.glob("python/rpa-script-*/rpa_script/__main__.py"))
    facts = cast("dict[str, object]", read(result, FACTS))
    assert facts["path"] == ""  # 運行環境的子程序拿不到電腦上的 PYTHONPATH
    workspace = Path(str(facts["cwd"]))
    assert workspace.name.startswith(".work-")
    assert not workspace.exists()  # 單獨執行：結束就刪
    steps = steps_of(result)
    assert [(s["id"], s["occurrence"], s["status"]) for s in steps] == [
        ("write", 1, "passed"),
        ("check", 1, "passed"),
        ("check", 2, "passed"),
    ]
    assert steps[0]["source"] == "automation.py:9"
    assert steps[0]["previews"] == ["previews/001.json"]
    preview = json.loads((result.directory / "previews" / "001.json").read_text("utf-8"))
    assert preview["columns"] == ["name", "value"]
    assert preview["rows"] == [["A", 3], ["B", 4]]
    assert (result.directory / "outputs" / "summary.csv").is_file()
    log = (result.directory / "automation.log").read_text(encoding="utf-8")
    assert "token ***" in log
    assert "secret-value-123" not in log
    assert not list((result.directory / "scenario").rglob("__pycache__"))
    saved = read(result, RESULT)
    assert isinstance(saved, dict)
    assert saved["python"] == result.python
    assert saved["kind"] == "script"


def test_nodes_of_one_composite_share_a_workspace(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    first = node(
        tmp_path / "first",
        "from pathlib import Path\nfrom rpa_script import automation\n"
        "@automation\ndef main(run):\n"
        "    (run.workspace / 'downloads').mkdir()\n"
        "    (run.workspace / 'downloads' / 'a.txt').write_text('甲', encoding='utf-8')\n",
        name="first",
    )
    second = node(
        tmp_path / "second",
        "from pathlib import Path\nfrom rpa_script import automation\n"
        "@automation\ndef main(run):\n"
        "    run.facts['seen'] = Path('downloads/a.txt').read_text(encoding='utf-8')\n",
        name="second",
    )
    one = execute(first, {}, out_root=tmp_path / "runs", workspace=workspace)
    two = execute(second, {}, out_root=tmp_path / "runs", workspace=workspace)
    assert one.status == two.status == "passed", (one.error, two.error)
    assert read(two, FACTS) == {"seen": "甲"}
    assert (workspace / "downloads" / "a.txt").is_file()  # 組合運行結束前不刪


def test_step_error_shows_only_the_message(tmp_path: Path) -> None:
    path = node(
        tmp_path / "scene",
        "from rpa_script import StepError, automation\n"
        "@automation\ndef main(run):\n"
        "    with run.step('find', '找報表'):\n"
        "        raise StepError('沒有 2026-10 的報表')\n",
    )
    result = execute(path, {}, out_root=tmp_path / "runs")
    assert result.status == "failed"
    assert result.error is not None
    assert "找報表：沒有 2026-10 的報表" in result.error
    log = (result.directory / "automation.log").read_text(encoding="utf-8")
    assert "Traceback" not in log
    assert not (result.directory / OUTPUT).exists()


def test_error_outside_steps_is_recorded_with_its_line(tmp_path: Path) -> None:
    path = node(tmp_path / "scene", "import json\n\nvalue = 1 / 0\n")
    result = execute(path, {}, out_root=tmp_path / "runs")
    assert result.status == "failed"
    steps = steps_of(result)
    assert steps[-1]["id"] == "script"
    assert steps[-1]["source"] == "automation.py:3"
    assert "ZeroDivisionError" in str(steps[-1]["error"])
    log = (result.directory / "automation.log").read_text(encoding="utf-8")
    assert "automation.py" in log
    assert "rpa_script" not in log  # 呼叫堆疊只留場景自己的檔案


def test_plain_scripts_keep_the_file_contract(tmp_path: Path) -> None:
    path = node(
        tmp_path / "scene",
        "import argparse, json\nfrom pathlib import Path\n"
        "parser = argparse.ArgumentParser()\n"
        "for name in ('input', 'facts', 'screenshot', 'target-url'):\n"
        "    parser.add_argument(f'--{name}', required=True)\n"
        "args = parser.parse_args()\n"
        "params = json.loads(Path(args.input).read_text(encoding='utf-8'))\n"
        "Path(args.facts).write_text(json.dumps({'echo': params['x'], 'url': args.target_url}),"
        " encoding='utf-8')\n",
        extra="inputs:\n  x: { type: object }\n",
    )
    result = execute(path, {"x": {"a": [1, 2]}}, out_root=tmp_path / "runs")
    assert result.status == "passed", result.error
    assert read(result, FACTS) == {"echo": {"a": [1, 2]}, "url": ""}


def test_assertion_runs_in_the_runtime(tmp_path: Path, runtime_root: Path) -> None:
    path = node(
        tmp_path / "scene",
        "from rpa_script import automation\n@automation\ndef main(run):\n    run.facts['n'] = 3\n",
        extra="outputs:\n  n: { type: integer }\n  exe: { type: string }\n",
        assertion=(
            "import argparse, json, sys\nfrom pathlib import Path\n"
            "p = argparse.ArgumentParser()\n"
            "for n in ('input', 'facts', 'output'):\n    p.add_argument(f'--{n}', required=True)\n"
            "a = p.parse_args()\n"
            "facts = json.loads(Path(a.facts).read_text(encoding='utf-8'))\n"
            "Path(a.output).write_text(json.dumps({'n': facts['n'], 'exe': sys.executable}),"
            " encoding='utf-8')\n"
        ),
    )
    result = execute(path, {}, out_root=tmp_path / "runs")
    assert result.status == "passed", (result.error, result.problems)
    assert result.output is not None
    assert result.output["n"] == 3
    assert str(result.output["exe"]).startswith(str(runtime_root))  # 斷言也用運行環境的 Python


def test_missing_runtime_stops_before_the_script(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv(ENV_RUNTIME)
    monkeypatch.setenv("RPA_CACHE_DIR", str(tmp_path / "cache"))
    monkeypatch.setenv("RPA_PYTHON_ARCHIVES", str(tmp_path / "none"))
    path = node(tmp_path / "scene", "raise SystemExit('should not run')\n")
    result = execute(path, {}, out_root=tmp_path / "runs")
    assert result.status == "failed"
    assert result.stage("runtime").status == "failed"
    assert "RPA Python 運行環境 1 準備失敗" in str(result.error)
    assert result.stage("automation").status == "skipped"


def test_platforms_other_than_this_computer(tmp_path: Path) -> None:
    other = {"windows": "linux", "macos": "linux", "linux": "windows"}[platform_tag().split("-")[0]]
    path = node(tmp_path / "scene", "print('no')\n")
    path.write_text(
        path.read_text(encoding="utf-8") + f"  platforms: [{other}]\n", encoding="utf-8"
    )
    result = execute(path, {}, out_root=tmp_path / "runs")
    assert result.status == "failed"
    assert f"只在 {other} 執行" in str(result.error)


def test_deadline_stops_the_script_and_its_children(tmp_path: Path) -> None:
    path = node(
        tmp_path / "scene",
        "import subprocess, sys, time\n"
        "subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'])\n"
        "time.sleep(60)\n",
        extra="deadline: 1500\n",
    )
    started = time.monotonic()
    result = execute(path, {}, out_root=tmp_path / "runs")
    assert result.status == "timed_out"
    assert time.monotonic() - started < 15


def test_file_inputs_reach_the_script(tmp_path: Path) -> None:
    source = tmp_path / "報表.csv"
    source.write_text("a,b\n1,2\n", encoding="utf-8")
    path = node(
        tmp_path / "scene",
        "from pathlib import Path\nfrom rpa_script import automation\n"
        "@automation\ndef main(run):\n"
        "    files = run.input_files['report']\n"
        "    run.facts['lines'] = len(Path(files[0]).read_text(encoding='utf-8').splitlines())\n",
        extra="inputs:\n  report: { type: file }\n",
    )
    result = execute(
        path, {"report": {"path": str(source)}}, out_root=tmp_path / "runs", any_path=True
    )
    assert result.status == "passed", (result.error, result.problems)
    assert read(result, FACTS) == {"lines": 2}


EXCEL = """\
import pandas as pd
from openpyxl import load_workbook
from rpa_script import automation


@automation
def main(run):
    with run.step("merge", "合併分店報表"):
        frames = [pd.DataFrame(rows, columns=["品項", "金額"]).assign(分店=name)
                  for name, rows in run.inputs["branches"].items()]
        merged = pd.concat(frames, ignore_index=True)
        target = run.output_dir / "合併.xlsx"
        merged.to_excel(target, index=False, engine="xlsxwriter")
        run.preview("合併結果", merged)
    with run.step("check", "讀回檢查"):
        sheet = load_workbook(target).active
        run.facts["rows"] = sheet.max_row - 1
        run.facts["total"] = int(merged["金額"].sum())
"""


@real_only
def test_runtime_packages_make_an_excel_report(tmp_path: Path) -> None:
    path = node(
        tmp_path / "scene",
        EXCEL,
        extra=(
            "inputs:\n  branches: { type: object }\n"
            "outputs:\n  rows: { type: integer, from: '{{ facts.rows }}' }\n"
            "  total: { type: integer, from: '{{ facts.total }}' }\n"
            "verify:\n  - { value: '{{ facts.rows }}', equals: '3' }\n"
        ),
    )
    branches = {"台北": [["筆電", 100], ["滑鼠", 5]], "高雄": [["螢幕", 30]]}
    result = execute(path, {"branches": branches}, out_root=tmp_path / "runs")
    assert result.status == "passed", (result.error, result.problems)
    assert result.output == {"rows": 3, "total": 135}
    assert (result.directory / "outputs" / "合併.xlsx").is_file()
    preview = json.loads((result.directory / "previews" / "001.json").read_text("utf-8"))
    assert preview["columns"] == ["品項", "金額", "分店"]


@real_only
def test_real_runtime_is_isolated(tmp_path: Path) -> None:
    path = node(
        tmp_path / "scene",
        "import importlib.util, sys\nfrom rpa_script import automation\n"
        "@automation\ndef main(run):\n"
        "    run.facts['pip'] = importlib.util.find_spec('pip') is not None\n"
        "    run.facts['rpa_core'] = importlib.util.find_spec('rpa_core') is not None\n"
        "    run.facts['flags'] = [sys.flags.ignore_environment, sys.flags.no_user_site]\n",
    )
    result = execute(path, {}, out_root=tmp_path / "runs")
    assert result.status == "passed", (result.error, result.problems)
    assert read(result, FACTS) == {"pip": False, "rpa_core": False, "flags": [1, 1]}


# ---------------------------------------------------------------- 產出的檔案（file 出參與檔案變化）

SUMMARY = """\
from rpa_script import automation


@automation
def main(run):
    with run.step("write", "寫出彙總"):
        target = run.output_dir / "彙總.csv"
        target.write_text("分店,金額\\n台北,10\\n", encoding="utf-8", newline="\\n")
        (run.workspace / "notes").mkdir(exist_ok=True)
        (run.workspace / "notes" / "log.txt").write_text("ok", encoding="utf-8")
    run.facts["summary"] = str(target)
"""
SUMMARY_EXTRA = (
    "outputs:\n  summary: { type: file, from: '{{ facts.summary }}' }\n"
    "verify:\n  - { value: '{{ facts.summary }}', matches: 'csv$' }\n"
)


def test_file_output_is_published_with_its_description(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    path = node(tmp_path / "scene", SUMMARY, extra=SUMMARY_EXTRA)
    result = execute(path, {}, out_root=tmp_path / "runs", workspace=workspace)

    assert result.status == "passed", (result.error, result.problems)
    assert result.output is not None
    value = cast("dict[str, object]", result.output["summary"])
    content = "分店,金額\n台北,10\n".encode()
    assert value["name"] == "彙總.csv"
    assert value["size"] == len(content)
    assert value["workspace"] == f".outputs/{result.run_id}/彙總.csv"
    assert (result.directory / "outputs" / "彙總.csv").read_bytes() == content
    assert (workspace / str(value["workspace"])).read_bytes() == content  # 下游節點拿得到
    changes = cast("dict[str, object]", read(result, "files.json"))
    added = cast("list[dict[str, object]]", changes["added"])
    assert {(item["where"], item["path"]) for item in added} == {
        ("outputs", "彙總.csv"),
        ("workspace", "notes/log.txt"),
    }
    assert changes["truncated"] is False
    # 每一步前後的變化（rpa_script）
    step = steps_of(result)[0]
    assert step["files"] == {"added": ["outputs/彙總.csv", "workspace/notes/log.txt"]}
    report = (result.directory / "report.html").read_text(encoding="utf-8")
    assert "產出檔案" in report
    assert 'href="outputs/%E5%BD%99%E7%B8%BD.csv"' in report
    assert "檔案變化" in report
    assert "notes/log.txt" in report
    assert "檔案：新增 2、改變 0、刪除 0" in report


def test_file_outputs_pass_through_the_shared_workspace(tmp_path: Path) -> None:
    """上游的 file 出參直接當下游的 file 入參（同一次組合運行、同一個共用工作資料夾）。"""
    workspace = tmp_path / "workspace"
    first = node(
        tmp_path / "first",
        "from rpa_script import automation\n@automation\ndef main(run):\n"
        "    (run.workspace / 'data').mkdir()\n"
        "    (run.workspace / 'data' / '明細.txt').write_text('甲乙丙', encoding='utf-8')\n"
        "    run.facts['data'] = 'data/明細.txt'\n",
        extra="outputs:\n  data: { type: file, from: '{{ facts.data }}' }\n"
        "verify:\n  - { value: '{{ facts.data }}', notEmpty: true }\n",
        name="first",
    )
    second = node(
        tmp_path / "second",
        "from rpa_script import automation\n@automation\ndef main(run):\n"
        "    [path] = run.input_files['data']\n"
        "    run.facts['text'] = path.read_text(encoding='utf-8')\n"
        "    run.facts['name'] = run.inputs['data']['name']\n",
        extra="inputs:\n  data: { type: file }\n",
        name="second",
    )
    one = execute(first, {}, out_root=tmp_path / "runs", workspace=workspace)
    assert one.status == "passed", (one.error, one.problems)
    assert one.output is not None
    value = one.output["data"]
    assert cast("dict[str, object]", value)["workspace"] == "data/明細.txt"
    two = execute(second, {"data": value}, out_root=tmp_path / "runs", workspace=workspace)
    assert two.status == "passed", (two.error, two.problems)
    assert read(two, FACTS) == {"text": "甲乙丙", "name": "明細.txt"}
    # 單獨執行沒有共用工作資料夾：說明這個檔案只有組合運行裡才拿得到
    alone = execute(second, {"data": value}, out_root=tmp_path / "runs")
    assert alone.status == "failed"
    assert "只有在組合運行裡才拿得到" in " ".join(alone.problems or [])


@pytest.mark.parametrize(
    ("fact", "message"),
    [
        ("'/'", "要在共用工作資料夾或這次執行的 outputs/ 裡"),
        ("'missing.csv'", "不存在"),
        ("['a.csv', 'b.csv']", "只能是一個檔案"),
    ],
)
def test_bad_file_outputs_publish_nothing(tmp_path: Path, fact: str, message: str) -> None:
    path = node(
        tmp_path / "scene",
        f"from rpa_script import automation\n@automation\ndef main(run):\n"
        f"    run.facts['f'] = {fact}\n",
        extra="outputs:\n  f: { type: file, from: '{{ facts.f }}' }\n"
        "verify:\n  - { value: '{{ facts.f }}', notEmpty: true }\n",
    )
    result = execute(path, {}, out_root=tmp_path / "runs")
    assert result.status == "failed"
    assert result.stage("output").status == "failed"
    assert message in " ".join(result.problems or [])
    assert not (result.directory / OUTPUT).exists()


def test_relative_out_folder_works_from_the_workspace(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """rpa run 預設 --out runs（相對路徑）：腳本在共用工作資料夾執行，仍找得到快照裡的腳本。"""
    script = (
        "from rpa_script import automation\n@automation\ndef main(run):\n    run.facts['ok'] = 1\n"
    )
    path = node(tmp_path / "scene", script)
    monkeypatch.chdir(tmp_path)
    result = execute(path, {}, out_root=Path("runs"))
    assert result.status == "passed", (result.error, result.problems)
    assert result.directory.is_absolute()
