"""合併與寫出 Excel 的小工具（automation.py 用；放在 script.files 裡，跟著場景的版本走）。"""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SHEET = "合併"


def merge(branches: dict[str, list[list[object]]]) -> pd.DataFrame:
    frames = [
        pd.DataFrame(rows, columns=["品項", "金額"]).assign(分店=name)
        for name, rows in branches.items()
    ]
    return pd.concat(frames, ignore_index=True)


def write_excel(frame: pd.DataFrame, target: Path) -> str:
    frame.to_excel(target, index=False, sheet_name=SHEET, engine="xlsxwriter")
    sheet = load_workbook(target).active
    if sheet is None or sheet.max_row != len(frame) + 1:
        raise ValueError(f"{target.name} 讀回來的列數不對")
    return sheet.title
