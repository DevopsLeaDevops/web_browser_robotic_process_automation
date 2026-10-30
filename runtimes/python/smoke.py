"""RPA Python 運行環境的自我測試（ADR 0020）：建置後用剛建好的直譯器執行，全部通過才打包。

    <運行環境>/bin/python3 -E -s -X utf8 runtimes/python/smoke.py \\
        --rpa-script packages/rpa_script/src

（Windows 的直譯器是運行環境根目錄的 python.exe。）

在暫存資料夾（中文路徑）裡實際做一次每種典型的本機自動化：讀寫 xlsx、合併多個 Excel、
Excel 與 CSV（Big5）互轉、Word 範本填值、合併 PDF、壓縮；最後確認隔離設定與 rpa_script
（運行環境裡沒有它：和 Runner 一樣把 ``--rpa-script`` 的資料夾加在 import 路徑最後）。
只印一行一個結果，任何一項失敗就以結束碼 1 結束。
"""

import argparse
import csv
import importlib
import io
import json
import sys
import tempfile
import zipfile
from collections.abc import Callable
from pathlib import Path

MODULES = (
    "openpyxl",
    "xlsxwriter",
    "xlrd",
    "pandas",
    "numpy",
    "dateutil",
    "docx",
    "pypdf",
    "charset_normalizer",
    "requests",
    "sqlite3",
    "ssl",
    "zoneinfo",
)


def check_imports(_: Path) -> str:
    versions = []
    for name in MODULES:
        module = importlib.import_module(name)
        versions.append(f"{name} {getattr(module, '__version__', '')}".strip())
    return ", ".join(versions)


def check_isolation(_: Path) -> str:
    flags = sys.flags
    if not (flags.ignore_environment and flags.no_user_site):
        raise AssertionError("要用 -E -s 執行（不讀 PYTHON* 環境變數與使用者的 site-packages）")
    for name in ("pip", "rpa_script", "rpa_core"):
        try:
            importlib.import_module(name)
        except ImportError:
            continue
        raise AssertionError(f"運行環境裡不應該有 {name}")
    return sys.version.split()[0]


def check_excel(folder: Path) -> str:
    import openpyxl
    import pandas as pd

    sources = []
    for index, branch in enumerate(("台北", "高雄", "台中"), start=1):
        book = openpyxl.Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "銷售"
        sheet.append(["品項", "金額"])
        for row in range(500):
            sheet.append([f"{branch}-{row}", index * 10 + row % 7])
        path = folder / f"分店報表-{branch}.xlsx"
        book.save(path)
        sources.append(path)
    frames = [pd.read_excel(path, sheet_name="銷售").assign(分店=path.stem) for path in sources]
    merged = pd.concat(frames, ignore_index=True)
    target = folder / "合併.xlsx"
    merged.to_excel(target, index=False, engine="xlsxwriter")
    again = pd.read_excel(target)
    if len(again) != 1500 or int(again["金額"].sum()) != int(merged["金額"].sum()):
        raise AssertionError("合併後的列數或總額不對")
    return f"{len(again)} 列，總額 {int(again['金額'].sum())}"


def check_csv_big5(folder: Path) -> str:
    import pandas as pd
    from charset_normalizer import from_path

    path = folder / "舊系統匯出.csv"
    with path.open("w", encoding="big5", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["客戶", "地址"])
        for index in range(50):
            writer.writerow([f"客戶{index}", "臺北市中正區重慶南路一段"])
    best = from_path(path).best()
    if best is None:
        raise AssertionError("判斷不出編碼")
    frame = pd.read_csv(path, encoding=best.encoding)
    if frame.iloc[0]["地址"] != "臺北市中正區重慶南路一段":
        raise AssertionError(f"用 {best.encoding} 讀出來的內容不對")
    frame.to_excel(folder / "舊系統匯出.xlsx", index=False)
    return best.encoding


def check_word(folder: Path) -> str:
    import docx

    template = docx.Document()
    template.add_paragraph("親愛的 {{客戶}}：本月金額 {{金額}} 元。")
    path = folder / "範本.docx"
    template.save(path)
    filled = docx.Document(str(path))
    for paragraph in filled.paragraphs:
        paragraph.text = paragraph.text.replace("{{客戶}}", "王小明").replace("{{金額}}", "1,200")
    out = folder / "通知.docx"
    filled.save(out)
    text = docx.Document(str(out)).paragraphs[0].text
    if "王小明" not in text:
        raise AssertionError("Word 範本沒有填進去")
    return text


def check_pdf(folder: Path) -> str:
    from pypdf import PdfReader, PdfWriter

    parts = []
    for count in (1, 2):
        writer = PdfWriter()
        for _ in range(count):
            writer.add_blank_page(width=595, height=842)
        path = folder / f"第{count}份.pdf"
        with path.open("wb") as handle:
            writer.write(handle)
        parts.append(path)
    merged = PdfWriter()
    for path in parts:
        merged.append(str(path))
    target = folder / "合併.pdf"
    with target.open("wb") as handle:
        merged.write(handle)
    pages = len(PdfReader(str(target)).pages)
    if pages != 3:
        raise AssertionError(f"合併後應該有 3 頁，實際 {pages} 頁")
    return f"{pages} 頁"


def check_zip(folder: Path) -> str:
    target = folder / "報表.zip"
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(folder.glob("*.xlsx")):
            archive.write(path, path.name)
    with zipfile.ZipFile(target) as archive:
        names = archive.namelist()
    if "合併.xlsx" not in names:
        raise AssertionError("壓縮檔裡沒有中文檔名的檔案")
    return f"{len(names)} 個檔案"


def check_rpa_script(_: Path) -> str:
    import rpa_script

    return json.dumps(rpa_script.selfcheck(), ensure_ascii=False)


CHECKS: tuple[tuple[str, Callable[[Path], str]], ...] = (
    ("隔離", check_isolation),
    ("import", check_imports),
    ("Excel", check_excel),
    ("CSV（Big5）", check_csv_big5),
    ("Word", check_word),
    ("PDF", check_pdf),
    ("壓縮", check_zip),
    ("rpa_script", check_rpa_script),
)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rpa-script", type=Path, required=True, help="rpa_script 所在的資料夾")
    args = parser.parse_args()
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", line_buffering=True)
    failed = 0
    with tempfile.TemporaryDirectory(prefix="rpa-運行環境-") as temp:
        folder = Path(temp) / "中文 資料夾"
        folder.mkdir()
        for name, check in CHECKS:
            if check is check_rpa_script:
                sys.path.append(str(args.rpa_script.resolve()))  # 和 Runner 的啟動方式相同
            try:
                detail = check(folder)
            except Exception as error:  # 每一項都要跑完，最後一起說明
                failed += 1
                sys.stdout.write(f"✗ {name}：{type(error).__name__}: {error}\n")
            else:
                sys.stdout.write(f"✓ {name}：{detail}\n")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
